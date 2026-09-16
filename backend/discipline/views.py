import io
from django.db.models import Q
from django.http import HttpResponse
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.graphics.barcode import code128
from reportlab.lib.units import mm
from PIL import Image as PILImage

from django.utils import timezone
from .models import (
    DisciplineApprovalFlowConfig,
    DisciplineCategory,
    DisciplineIncidentLog,
    DisciplineIncidentAction,
)
from .serializers import (
    DisciplineApprovalFlowConfigSerializer,
    DisciplineCategorySerializer,
    DisciplineIncidentLogSerializer,
    DisciplineIncidentActionSerializer,
)
from academics.models import (
    StudentProfile,
    Section,
    Batch,
    Department,
    AcademicYear,
    StudentMentorMap,
    SectionAdvisor,
    DepartmentRole,
    StaffProfile,
)
from accounts.models import User, Role
from academics.services import authority_resolver


def is_user_approver_for_incident(user, log: DisciplineIncidentLog, target_role: str = None) -> bool:
    """
    Check if `user` is the specific assigned approver for `log` when `target_role` is required.
    Handles semantic roles (HOD, AHOD, MENTOR, ADVISOR) by mapping to the specific student's
    mentor, section advisor, or department HOD/AHOD.
    Also handles system administrative roles (IQAC, ADMIN, DISCIPLINE_COMMITTEE_ADMIN) and
    direct role assignments.
    """
    if not user or not user.is_authenticated:
        return False

    role_code = (target_role or log.current_role or '').strip().upper()
    if not role_code:
        return False

    user_roles = list(user.roles.values_list('name', flat=True))
    if hasattr(user, 'role') and user.role:
        user_roles.append(user.role)
    roles_upper = {r.upper() for r in user_roles}

    # Administrative bypass: Admin and IQAC have global authority
    if any(r in ['IQAC', 'ADMIN', 'DISCIPLINE_COMMITTEE_ADMIN', 'DISCIPLINECOMMITTEEADMIN'] for r in roles_upper):
        return True

    # Find the StudentProfile for this incident log
    student = log.student
    if student is None and log.reg_no:
        student = StudentProfile.objects.filter(
            Q(reg_no__iexact=log.reg_no) | Q(user__username__iexact=log.reg_no)
        ).select_related(
            'user', 'home_department', 'section', 'section__batch',
            'section__batch__course', 'section__batch__course__department'
        ).first()

    academic_year = AcademicYear.objects.filter(is_active=True).first() or AcademicYear.objects.order_by('-id').first()

    # 1. MENTOR role check: Must be the specific student's assigned active mentor
    if role_code == 'MENTOR':
        if student:
            mentor_staff = authority_resolver.get_student_mentor(student)
            if mentor_staff and mentor_staff.user_id == user.id:
                return True
        return False

    # 2. ADVISOR role check: Must be the specific student's section advisor
    if role_code == 'ADVISOR':
        if student and academic_year:
            advisor_staff = authority_resolver.get_section_advisor(student, academic_year)
            if advisor_staff and advisor_staff.user_id == user.id:
                return True
        return False

    # 3. HOD role check: Must be the student's department HOD
    if role_code == 'HOD':
        dept = None
        if student:
            dept = authority_resolver._get_student_department(student)
        if not dept and log.department_name:
            dept = Department.objects.filter(
                Q(name__iexact=log.department_name) | Q(short_name__iexact=log.department_name)
            ).first()

        if dept and academic_year:
            hod_staff = authority_resolver.get_department_hod_by_department(dept, academic_year)
            if hod_staff and hod_staff.user_id == user.id:
                return True
        return False

    # 4. AHOD role check: Must be the student's department AHOD
    if role_code == 'AHOD':
        dept = None
        if student:
            dept = authority_resolver._get_student_department(student)
        if not dept and log.department_name:
            dept = Department.objects.filter(
                Q(name__iexact=log.department_name) | Q(short_name__iexact=log.department_name)
            ).first()

        if dept and academic_year:
            ahod_staff = authority_resolver.get_department_ahod_by_department(dept, academic_year)
            if ahod_staff and ahod_staff.user_id == user.id:
                return True
        return False

    # 5. Generic / Committee role match (e.g. DISCIPLINE_COMMITTEE, PRINCIPAL, DEAN, etc.)
    # Match user's direct assigned roles
    clean_role_code = role_code.replace(' ', '_').replace('-', '_')
    for u_role in roles_upper:
        clean_u_role = u_role.replace(' ', '_').replace('-', '_')
        if clean_u_role == clean_role_code:
            return True

    return False


class DisciplineApprovalFlowConfigView(APIView):
    """
    GET /api/discipline/flow/ - Retrieve active approval flow config.
    POST /api/discipline/flow/ - Update or create approval flow config.
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request):
        flow = DisciplineApprovalFlowConfig.objects.filter(is_active=True).first()
        if not flow:
            # Default starter flow if none exists
            flow = DisciplineApprovalFlowConfig.objects.create(
                name='Standard Discipline Approval Flow',
                roles=['DISCIPLINE_COMMITTEE', 'HOD', 'PRINCIPAL'],
                is_active=True,
            )
        serializer = DisciplineApprovalFlowConfigSerializer(flow)

        # Also return list of all available system roles for convenience
        all_roles = list(Role.objects.values_list('name', flat=True).distinct())
        # Filter out student/ext_staff if needed or sort
        return Response({
            'flow': serializer.data,
            'available_roles': sorted(all_roles),
        })

    def post(self, request):
        roles = request.data.get('roles', [])
        if not isinstance(roles, list) or len(roles) == 0:
            return Response({'error': 'At least one approver role must be in the flow.'}, status=status.HTTP_400_BAD_REQUEST)

        name = request.data.get('name', 'Standard Discipline Approval Flow')
        flow = DisciplineApprovalFlowConfig.objects.filter(is_active=True).first()
        if flow:
            flow.name = name
            flow.roles = roles
            flow.save()
        else:
            flow = DisciplineApprovalFlowConfig.objects.create(name=name, roles=roles, is_active=True)

        serializer = DisciplineApprovalFlowConfigSerializer(flow)
        return Response(serializer.data, status=status.HTTP_200_OK)


class DisciplineCategoriesView(APIView):
    """
    GET /api/discipline/categories/ - List all configured discipline categories.
    POST /api/discipline/categories/ - Create a new discipline category.
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request):
        qs = DisciplineCategory.objects.all()
        is_active = request.query_params.get('is_active')
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ['1', 'true'])

        sem = request.query_params.get('semester')
        if sem and sem != 'ALL':
            qs = qs.filter(semesters__contains=[sem])

        serializer = DisciplineCategorySerializer(qs, many=True)
        return Response({'results': serializer.data})

    def post(self, request):
        serializer = DisciplineCategorySerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class DisciplineCategoryDetailView(APIView):
    """
    GET, PUT, PATCH, DELETE /api/discipline/categories/<id>/
    """
    permission_classes = (IsAuthenticated,)

    def get_object(self, pk):
        try:
            return DisciplineCategory.objects.get(pk=pk)
        except DisciplineCategory.DoesNotExist:
            return None

    def get(self, request, pk):
        category = self.get_object(pk)
        if not category:
            return Response({'error': 'Category not found'}, status=status.HTTP_404_NOT_FOUND)
        serializer = DisciplineCategorySerializer(category)
        return Response(serializer.data)

    def put(self, request, pk):
        category = self.get_object(pk)
        if not category:
            return Response({'error': 'Category not found'}, status=status.HTTP_404_NOT_FOUND)
        serializer = DisciplineCategorySerializer(category, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def patch(self, request, pk):
        return self.put(request, pk)

    def delete(self, request, pk):
        category = self.get_object(pk)
        if not category:
            return Response({'error': 'Category not found'}, status=status.HTTP_404_NOT_FOUND)
        category.delete()
        return Response({'success': True, 'message': 'Category deleted successfully'}, status=status.HTTP_200_OK)


class DisciplineIncidentLogsView(APIView):
    """
    GET /api/discipline/logs/ - List real incident logs with filters.
    POST /api/discipline/logs/ - Record a new student discipline incident.
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request):
        qs = DisciplineIncidentLog.objects.select_related('student', 'category', 'reported_by').prefetch_related('actions').all()

        search = request.query_params.get('search', '').strip()
        if search:
            qs = qs.filter(
                Q(reg_no__icontains=search) |
                Q(student_name__icontains=search) |
                Q(username__icontains=search) |
                Q(category_title__icontains=search) |
                Q(remarks__icontains=search)
            )

        status_param = request.query_params.get('status')
        if status_param and status_param != 'ALL':
            qs = qs.filter(status=status_param)

        severity_param = request.query_params.get('severity')
        if severity_param and severity_param != 'ALL':
            qs = qs.filter(severity=severity_param)

        dept_param = request.query_params.get('department')
        if dept_param and dept_param != 'ALL':
            qs = qs.filter(
                Q(department_name__iexact=dept_param) |
                Q(student__home_department__short_name__iexact=dept_param) |
                Q(student__home_department__name__iexact=dept_param) |
                Q(student__section__batch__course__department__short_name__iexact=dept_param) |
                Q(student__section__batch__course__department__name__iexact=dept_param)
            )

        date_param = request.query_params.get('date')
        if date_param:
            qs = qs.filter(incident_date__date=date_param)

        # Optional filter for logged in user's role approvals
        role_filter = request.query_params.get('current_role')
        if role_filter:
            qs = qs.filter(current_role__iexact=role_filter)

        total_count = qs.count()
        reported_count = qs.filter(status__in=['REPORTED', 'PENDING_FINE_PAYMENT', 'RECEIPT_UPLOADED', 'IN_APPROVAL']).count()
        action_taken_count = qs.filter(status='ACTION_TAKEN').count()
        resolved_count = qs.filter(status='RESOLVED').count()

        available_depts = list(Department.objects.values_list('short_name', flat=True).filter(short_name__isnull=False).distinct())
        dept_names = list(Department.objects.values_list('name', flat=True).filter(name__isnull=False).distinct())
        all_dept_options = sorted(list(set([d for d in (available_depts + dept_names) if d])))

        serializer = DisciplineIncidentLogSerializer(qs, many=True)
        return Response({
            'results': serializer.data,
            'departments': all_dept_options,
            'total_count': total_count,
            'reported_count': reported_count,
            'action_taken_count': action_taken_count,
            'resolved_count': resolved_count,
        })

    def post(self, request):
        data = request.data.copy()

        # Set reporting user details
        if request.user.is_authenticated:
            data['reported_by'] = request.user.id
            full_reporter_name = f"{request.user.first_name} {request.user.last_name}".strip()
            data['reported_by_name'] = full_reporter_name or request.user.username

        # Sanitize student ID if passed as 0, null, or string
        student_val = data.get('student')
        if not student_val or str(student_val) in ['0', 'null', 'None', '']:
            data['student'] = None
        else:
            try:
                data['student'] = int(student_val)
            except (ValueError, TypeError):
                data['student'] = None

        # If student is not explicitly linked, try to look up by register number or digits
        reg_no = str(data.get('reg_no', '')).strip()
        if reg_no:
            clean_digits = ''.join([c for c in reg_no if c.isdigit()])
            st_query = (
                Q(reg_no__iexact=reg_no) |
                Q(user__username__iexact=reg_no)
            )
            if clean_digits:
                st_query |= Q(reg_no__icontains=clean_digits)

            st = StudentProfile.objects.filter(st_query).select_related(
                'user', 'home_department', 'section', 'section__batch', 'section__batch__course', 'section__batch__course__department'
            ).first()

            if st:
                data['student'] = st.id
                if not data.get('student_name'):
                    st_user = st.user
                    full_name = f"{getattr(st_user, 'first_name', '')} {getattr(st_user, 'last_name', '')}".strip()
                    data['student_name'] = full_name or getattr(st_user, 'username', reg_no)
                data['username'] = getattr(st.user, 'username', '')
                if not data.get('department_name'):
                    data['department_name'] = st.home_department.name if st.home_department else (
                        st.section.batch.course.department.name if st.section and st.section.batch and st.section.batch.course and st.section.batch.course.department else ''
                    )
                if not data.get('section_name') and st.section:
                    data['section_name'] = st.section.name
                if not data.get('batch_name') and st.section and st.section.batch:
                    data['batch_name'] = st.section.batch.name

        cat_id = data.get('category')
        if cat_id:
            try:
                cat_obj = DisciplineCategory.objects.get(pk=cat_id)
                if not data.get('category_title'):
                    data['category_title'] = cat_obj.title
                if not data.get('severity'):
                    data['severity'] = cat_obj.severity
            except DisciplineCategory.DoesNotExist:
                pass

        # Attach active approval flow path to the incident
        active_flow = DisciplineApprovalFlowConfig.objects.filter(is_active=True).first()
        if active_flow and active_flow.roles:
            data['flow_roles'] = active_flow.roles
            data['current_step_index'] = 0
            data['current_role'] = active_flow.roles[0]
        else:
            default_roles = ['DISCIPLINE_COMMITTEE', 'HOD', 'PRINCIPAL']
            data['flow_roles'] = default_roles
            data['current_step_index'] = 0
            data['current_role'] = default_roles[0]

        data['status'] = 'REPORTED'

        serializer = DisciplineIncidentLogSerializer(data=data)
        if serializer.is_valid():
            obj = serializer.save()
            # Record first action in audit
            DisciplineIncidentAction.objects.create(
                incident=obj,
                action_type='REPORTED',
                role_performed=data.get('current_role', 'DISCIPLINE_COMMITTEE'),
                user=request.user,
                user_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
                comments=f"Incident recorded: {obj.category_title or 'Discipline violation'}",
            )
            return Response(DisciplineIncidentLogSerializer(obj).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class DisciplineIncidentLogDetailView(APIView):
    """
    GET, PUT, PATCH, DELETE /api/discipline/logs/<id>/
    """
    permission_classes = (IsAuthenticated,)

    def get_object(self, pk):
        try:
            return DisciplineIncidentLog.objects.select_related('student', 'category', 'reported_by').prefetch_related('actions').get(pk=pk)
        except DisciplineIncidentLog.DoesNotExist:
            return None

    def get(self, request, pk):
        log = self.get_object(pk)
        if not log:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)
        serializer = DisciplineIncidentLogSerializer(log)
        return Response(serializer.data)

    def put(self, request, pk):
        log = self.get_object(pk)
        if not log:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)
        serializer = DisciplineIncidentLogSerializer(log, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def patch(self, request, pk):
        return self.put(request, pk)

    def delete(self, request, pk):
        log = self.get_object(pk)
        if not log:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)
        log.delete()
        return Response({'success': True, 'message': 'Incident log deleted successfully'}, status=status.HTTP_200_OK)


class DisciplineApprovalsListView(APIView):
    """
    GET /api/discipline/approvals/
    List incident logs pending approval or actionable by the current user.
    Only shows incidents where the user is the assigned approver for the student
    (e.g., student's specific Mentor, Section Advisor, or Department HOD/AHOD),
    or has an authorized administrative role.
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request):
        user = request.user
        user_roles = list(user.roles.values_list('name', flat=True))
        if hasattr(user, 'role') and user.role:
            user_roles.append(user.role)
        roles_upper = [r.upper() for r in user_roles]

        is_admin_or_iqac = any(r in ['IQAC', 'ADMIN', 'DISCIPLINE_COMMITTEE_ADMIN', 'DISCIPLINECOMMITTEEADMIN'] for r in roles_upper)

        qs = DisciplineIncidentLog.objects.select_related(
            'student', 'student__user', 'student__home_department', 'student__section',
            'student__section__batch', 'student__section__batch__course', 'student__section__batch__course__department',
            'category', 'reported_by'
        ).prefetch_related('actions').exclude(status__in=['RESOLVED', 'DISMISSED'])

        if is_admin_or_iqac:
            results = qs
        else:
            # Candidate filter: only check logs where user can act for that specific student's role
            matching_logs = []
            for log in qs:
                if is_user_approver_for_incident(user, log):
                    matching_logs.append(log)
            results = matching_logs

        serializer = DisciplineIncidentLogSerializer(results, many=True)
        return Response({'results': serializer.data})


class DisciplineFixFineView(APIView):
    """
    POST /api/discipline/logs/<id>/fix-fine/
    Fix or modify the fine amount (in Rupees) for an incident.
    Supports switch: fine / no fine.
    """
    permission_classes = (IsAuthenticated,)

    def post(self, request, pk):
        try:
            log = DisciplineIncidentLog.objects.get(pk=pk)
        except DisciplineIncidentLog.DoesNotExist:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)

        if not is_user_approver_for_incident(request.user, log):
            return Response({'error': 'You are not the designated approver for this student at this stage.'}, status=status.HTTP_403_FORBIDDEN)

        has_fine = request.data.get('has_fine', True)
        fine_amount = request.data.get('fine_amount', 0)
        comments = request.data.get('comments', '')

        user_role = request.data.get('role') or log.current_role or 'APPROVER'

        is_initial_fix = not log.has_fine or log.fine_amount == 0

        log.has_fine = bool(has_fine)
        log.fine_amount = fine_amount if has_fine else 0.00
        log.fine_fixed_by = request.user
        log.fine_fixed_at = timezone.now()

        if has_fine and float(fine_amount) > 0:
            if log.status == 'REPORTED':
                log.status = 'PENDING_FINE_PAYMENT'
        else:
            if log.status == 'PENDING_FINE_PAYMENT':
                log.status = 'IN_APPROVAL'

        log.save()

        # Audit Action
        DisciplineIncidentAction.objects.create(
            incident=log,
            action_type='FINE_FIXED' if is_initial_fix else 'FINE_MODIFIED',
            role_performed=user_role,
            user=request.user,
            user_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
            comments=comments or (f"Fine fixed at ₹{fine_amount}" if has_fine else "Set to No Fine"),
            fine_amount=fine_amount if has_fine else 0.00,
        )

        return Response(DisciplineIncidentLogSerializer(log).data)


class DisciplineForwardActionView(APIView):
    """
    POST /api/discipline/logs/<id>/forward/
    Forward the incident request to the next role in the approval flow.
    """
    permission_classes = (IsAuthenticated,)

    def post(self, request, pk):
        try:
            log = DisciplineIncidentLog.objects.get(pk=pk)
        except DisciplineIncidentLog.DoesNotExist:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)

        if not is_user_approver_for_incident(request.user, log):
            return Response({'error': 'You are not the designated approver for this student at this stage.'}, status=status.HTTP_403_FORBIDDEN)

        # Enforce fine payment receipt requirement before allowing forward/approval
        if log.has_fine and float(log.fine_amount or 0) > 0 and not (log.receipt_document or log.receipt_document_url):
            return Response({'error': 'Cannot forward incident request: Student has not uploaded fee payment receipt proof.'}, status=status.HTTP_400_BAD_REQUEST)

        flow_roles = log.flow_roles or []
        current_idx = log.current_step_index

        if current_idx + 1 >= len(flow_roles):
            return Response({'error': 'Request is already at the final approver in the flow. Please use Approve instead.'}, status=status.HTTP_400_BAD_REQUEST)

        next_idx = current_idx + 1
        next_role = flow_roles[next_idx]

        log.current_step_index = next_idx
        log.current_role = next_role
        log.status = 'IN_APPROVAL'
        log.save()

        comments = request.data.get('comments', '')
        user_role = request.data.get('role') or flow_roles[current_idx]

        DisciplineIncidentAction.objects.create(
            incident=log,
            action_type='FORWARDED',
            role_performed=user_role,
            user=request.user,
            user_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
            comments=comments or f"Forwarded to {next_role}",
        )

        return Response(DisciplineIncidentLogSerializer(log).data)


class DisciplineFinalApproveView(APIView):
    """
    POST /api/discipline/logs/<id>/approve/
    Final approver approves the incident -> moves status to RESOLVED / Completed.
    """
    permission_classes = (IsAuthenticated,)

    def post(self, request, pk):
        try:
            log = DisciplineIncidentLog.objects.get(pk=pk)
        except DisciplineIncidentLog.DoesNotExist:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)

        if not is_user_approver_for_incident(request.user, log):
            return Response({'error': 'You are not the designated approver for this student at this stage.'}, status=status.HTTP_403_FORBIDDEN)

        # Enforce fine payment receipt requirement before allowing final approval
        if log.has_fine and float(log.fine_amount or 0) > 0 and not (log.receipt_document or log.receipt_document_url):
            return Response({'error': 'Cannot approve incident request: Student has not uploaded fee payment receipt proof.'}, status=status.HTTP_400_BAD_REQUEST)

        comments = request.data.get('comments', '')
        action_notes = request.data.get('action_notes', '')

        log.status = 'RESOLVED'
        if action_notes:
            log.action_notes = action_notes
        log.save()

        user_role = request.data.get('role') or log.current_role or 'FINAL_APPROVER'

        DisciplineIncidentAction.objects.create(
            incident=log,
            action_type='APPROVED',
            role_performed=user_role,
            user=request.user,
            user_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
            comments=comments or "Final approval granted. Discipline incident closed & resolved.",
        )

        return Response(DisciplineIncidentLogSerializer(log).data)


class DisciplineStudentUploadReceiptView(APIView):
    """
    POST /api/discipline/logs/<id>/upload-receipt/
    Student uploads receipt proof for fine payment.
    """
    permission_classes = (IsAuthenticated,)

    def post(self, request, pk):
        try:
            log = DisciplineIncidentLog.objects.get(pk=pk)
        except DisciplineIncidentLog.DoesNotExist:
            return Response({'error': 'Incident log not found'}, status=status.HTTP_404_NOT_FOUND)

        file_obj = request.FILES.get('receipt_document')
        notes = request.data.get('receipt_notes', '')

        if file_obj:
            log.receipt_document = file_obj
            log.receipt_document_url = f"/media/discipline/receipts/{file_obj.name}"

        log.receipt_notes = notes
        log.receipt_uploaded_at = timezone.now()
        log.status = 'RECEIPT_UPLOADED'
        log.save()

        doc_url = log.receipt_document.url if log.receipt_document else log.receipt_document_url

        DisciplineIncidentAction.objects.create(
            incident=log,
            action_type='RECEIPT_UPLOADED',
            role_performed='STUDENT',
            user=request.user,
            user_name=f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
            comments=notes or "Fine payment receipt proof uploaded by student.",
            document_url=doc_url or '',
        )

        return Response(DisciplineIncidentLogSerializer(log).data)


class DisciplineStudentIncidentsView(APIView):
    """
    GET /api/discipline/student-my-incidents/
    Returns student's pending actions (where status != RESOLVED) and completed actions (RESOLVED).
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request):
        user = request.user
        st = StudentProfile.objects.filter(user=user).first()

        # Query incidents matching user id or student reg_no/username
        qs = DisciplineIncidentLog.objects.select_related('category', 'reported_by').prefetch_related('actions').filter(
            Q(student=st) if st else Q(username__iexact=user.username) | Q(reg_no__iexact=user.username)
        )

        pending = qs.exclude(status__in=['RESOLVED', 'DISMISSED'])
        completed = qs.filter(status__in=['RESOLVED', 'DISMISSED'])

        return Response({
            'pending_actions': DisciplineIncidentLogSerializer(pending, many=True).data,
            'completed_actions': DisciplineIncidentLogSerializer(completed, many=True).data,
        })


class DisciplineStudentDirectoryView(APIView):
    """
    GET /api/discipline/students/ - Stream and query real active students from DB with barcodes,
    avatars, batches, sections, and departments.
    """
    permission_classes = (IsAuthenticated,)

    def get(self, request):
        qs = StudentProfile.objects.select_related(
            'user', 'home_department', 'section', 'section__batch', 'section__batch__course', 'section__batch__course__department'
        ).filter(status__in=['ACTIVE', 'active', 'Active'])

        dept_param = request.query_params.get('department')
        if dept_param and dept_param != 'ALL':
            qs = qs.filter(
                Q(home_department__name__iexact=dept_param) |
                Q(home_department__short_name__iexact=dept_param) |
                Q(home_department__code__iexact=dept_param) |
                Q(section__batch__course__department__short_name__iexact=dept_param) |
                Q(section__batch__course__department__name__iexact=dept_param) |
                Q(section__batch__department__short_name__iexact=dept_param)
            )

        batch_param = request.query_params.get('batch')
        if batch_param and batch_param != 'ALL':
            qs = qs.filter(section__batch__name__iexact=batch_param)

        sec_param = request.query_params.get('section')
        if sec_param and sec_param != 'ALL':
            qs = qs.filter(section__name__iexact=sec_param)

        search = request.query_params.get('search', '').strip()
        if search:
            clean_digits = ''.join([c for c in search if c.isdigit()])
            search_filters = (
                Q(reg_no__icontains=search) |
                Q(user__username__icontains=search) |
                Q(user__first_name__icontains=search) |
                Q(user__last_name__icontains=search) |
                Q(user__email__icontains=search)
            )
            if clean_digits:
                search_filters |= Q(reg_no__icontains=clean_digits)
            qs = qs.filter(search_filters)

        results = []
        for st in qs:
            u = st.user
            full_name = f"{getattr(u, 'first_name', '')} {getattr(u, 'last_name', '')}".strip()

            dept = st.home_department
            if not dept and st.section and st.section.batch:
                course = getattr(st.section.batch, 'course', None)
                dept = (getattr(course, 'department', None) if course else None) or getattr(st.section.batch, 'department', None)

            dept_name = getattr(dept, 'short_name', None) or getattr(dept, 'code', None) or getattr(dept, 'name', '')
            batch_label = st.batch or (st.section.batch.name if st.section and st.section.batch else '')

            img_url = None
            if getattr(st, 'profile_image', None):
                try:
                    img_url = st.profile_image.url
                except Exception:
                    img_url = None
            elif getattr(u, 'profile_image', None):
                img_url = getattr(u, 'profile_image', None)

            results.append({
                'id': st.id,
                'reg_no': st.reg_no,
                'name': full_name or getattr(u, 'username', st.reg_no),
                'username': getattr(u, 'username', ''),
                'email': getattr(u, 'email', ''),
                'profile_image_url': img_url,
                'department': dept_name,
                'batch': batch_label,
                'section': getattr(st.section, 'name', '') if st.section else '',
                'status': getattr(st, 'status', 'ACTIVE'),
            })

        available_depts = list(Department.objects.values_list('short_name', flat=True).filter(short_name__isnull=False).distinct())
        available_batches = list(Batch.objects.values_list('name', flat=True).distinct())

        return Response({
            'count': len(results),
            'results': results,
            'departments': [d for d in available_depts if d],
            'batches': [b for b in available_batches if b],
        })


class DisciplineStudentExportExcelWithBarcodesView(APIView):
    """
    POST /api/discipline/students/export-excel/
    Generates an Excel spreadsheet containing rendered visual barcode images embedded directly in the Barcode column.
    Accepts student_ids (array) or filters in POST payload.
    """
    permission_classes = (IsAuthenticated,)

    def post(self, request):
        student_ids = request.data.get('student_ids', [])
        reg_nos = request.data.get('reg_nos', [])
        dept_param = request.data.get('department')
        batch_param = request.data.get('batch')
        sec_param = request.data.get('section')
        search = (request.data.get('search') or '').strip()

        qs = StudentProfile.objects.select_related(
            'user', 'home_department', 'section', 'section__batch', 'section__batch__course', 'section__batch__course__department'
        ).filter(status__in=['ACTIVE', 'active', 'Active'])

        if student_ids:
            qs = qs.filter(id__in=student_ids)
        elif reg_nos:
            qs = qs.filter(reg_no__in=reg_nos)
        else:
            if dept_param and dept_param != 'ALL':
                qs = qs.filter(
                    Q(home_department__name__iexact=dept_param) |
                    Q(home_department__short_name__iexact=dept_param) |
                    Q(home_department__code__iexact=dept_param) |
                    Q(section__batch__course__department__short_name__iexact=dept_param) |
                    Q(section__batch__course__department__name__iexact=dept_param) |
                    Q(section__batch__department__short_name__iexact=dept_param)
                )
            if batch_param and batch_param != 'ALL':
                qs = qs.filter(section__batch__name__iexact=batch_param)
            if sec_param and sec_param != 'ALL':
                qs = qs.filter(section__name__iexact=sec_param)
            if search:
                qs = qs.filter(
                    Q(reg_no__icontains=search) |
                    Q(user__username__icontains=search) |
                    Q(user__first_name__icontains=search) |
                    Q(user__last_name__icontains=search)
                )

        # Build Excel Workbook with openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "StudentBarcodes"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Arial", size=10)
        bold_font = Font(name="Arial", size=10, bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        border_thin = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        headers = [
            "S.No", "Student Name", "Username", "Register Number",
            "Department", "Batch", "Section", "Barcode Digits", "Barcode Image (CODE128)"
        ]

        ws.row_dimensions[1].height = 28
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Set Column Widths
        ws.column_dimensions['A'].width = 8   # S.No
        ws.column_dimensions['B'].width = 28  # Name
        ws.column_dimensions['C'].width = 18  # Username
        ws.column_dimensions['D'].width = 22  # Reg No
        ws.column_dimensions['E'].width = 18  # Dept
        ws.column_dimensions['F'].width = 14  # Batch
        ws.column_dimensions['G'].width = 10  # Section
        ws.column_dimensions['H'].width = 18  # Barcode Digits
        ws.column_dimensions['I'].width = 30  # Barcode Image column

        # Populate student rows and generate/insert barcode images
        row_num = 2
        for idx, st in enumerate(qs, 1):
            u = st.user
            full_name = f"{getattr(u, 'first_name', '')} {getattr(u, 'last_name', '')}".strip() or getattr(u, 'username', st.reg_no)
            dept = st.home_department
            if not dept and st.section and st.section.batch:
                course = getattr(st.section.batch, 'course', None)
                dept = (getattr(course, 'department', None) if course else None) or getattr(st.section.batch, 'department', None)

            dept_name = getattr(dept, 'short_name', None) or getattr(dept, 'code', None) or getattr(dept, 'name', '')
            batch_label = st.batch or (st.section.batch.name if st.section and st.section.batch else '')
            clean_digits = ''.join([c for c in str(st.reg_no) if c.isdigit()]) or str(st.reg_no)

            ws.row_dimensions[row_num].height = 42

            c_sno = ws.cell(row=row_num, column=1, value=idx)
            c_sno.alignment = center_align
            c_sno.font = regular_font
            c_sno.border = border_thin

            c_name = ws.cell(row=row_num, column=2, value=full_name)
            c_name.alignment = left_align
            c_name.font = bold_font
            c_name.border = border_thin

            c_user = ws.cell(row=row_num, column=3, value=getattr(u, 'username', ''))
            c_user.alignment = left_align
            c_user.font = regular_font
            c_user.border = border_thin

            c_reg = ws.cell(row=row_num, column=4, value=st.reg_no)
            c_reg.alignment = left_align
            c_reg.font = bold_font
            c_reg.border = border_thin

            c_dept = ws.cell(row=row_num, column=5, value=dept_name or '-')
            c_dept.alignment = left_align
            c_dept.font = regular_font
            c_dept.border = border_thin

            c_batch = ws.cell(row=row_num, column=6, value=batch_label or '-')
            c_batch.alignment = center_align
            c_batch.font = regular_font
            c_batch.border = border_thin

            c_sec = ws.cell(row=row_num, column=7, value=getattr(st.section, 'name', '') or '-')
            c_sec.alignment = center_align
            c_sec.font = regular_font
            c_sec.border = border_thin

            c_digits = ws.cell(row=row_num, column=8, value=clean_digits)
            c_digits.alignment = center_align
            c_digits.font = bold_font
            c_digits.border = border_thin

            c_img_cell = ws.cell(row=row_num, column=9)
            c_img_cell.border = border_thin

            # Generate barcode image using reportlab & Pillow and embed directly into cell
            try:
                barcode_drawing = code128.Code128(clean_digits, barHeight=12 * mm, barWidth=0.35 * mm, humanReadable=True)
                img_buffer = io.BytesIO()
                # Render reportlab drawing to PNG
                from reportlab.graphics import renderPM
                renderPM.drawToFile(barcode_drawing, img_buffer, fmt="PNG", dpi=150)
                img_buffer.seek(0)

                img = OpenpyxlImage(img_buffer)
                img.width = 175
                img.height = 48

                # Position the image in Column I of current row
                cell_coord = f"I{row_num}"
                ws.add_image(img, cell_coord)
            except Exception as e:
                c_img_cell.value = clean_digits

            row_num += 1

        output_stream = io.BytesIO()
        wb.save(output_stream)
        output_stream.seek(0)

        response = HttpResponse(
            output_stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="student_barcodes.xlsx"'
        return response


class DisciplineStudentExportBarcodeImagesZipView(APIView):
    """
    POST /api/discipline/students/export-zip/
    Generate and return a ZIP file containing separate, cropped barcode PNG images
    for all filtered or selected students.
    File name of each image: <register_number>.png
    Image content: ONLY the barcode bars (no register number text below/around).
    """
    permission_classes = (IsAuthenticated,)

    def post(self, request):
        import zipfile
        import re

        student_ids = request.data.get('student_ids', [])
        reg_nos = request.data.get('reg_nos', [])
        dept_param = request.data.get('department')
        batch_param = request.data.get('batch')
        sec_param = request.data.get('section')
        search = (request.data.get('search') or '').strip()

        qs = StudentProfile.objects.select_related(
            'user', 'home_department', 'section', 'section__batch', 'section__batch__course', 'section__batch__course__department'
        ).filter(status__in=['ACTIVE', 'active', 'Active'])

        if student_ids and len(student_ids) > 0:
            # Handle list of integer/string IDs
            id_list = []
            for s in student_ids:
                try:
                    id_list.append(int(s))
                except (ValueError, TypeError):
                    pass
            if id_list:
                qs = qs.filter(id__in=id_list)
            else:
                qs = qs.filter(reg_no__in=[str(r).strip() for r in student_ids])
        elif reg_nos and len(reg_nos) > 0:
            qs = qs.filter(reg_no__in=[str(r).strip() for r in reg_nos])
        else:
            if dept_param and dept_param != 'ALL':
                qs = qs.filter(
                    Q(home_department__name__iexact=dept_param) |
                    Q(home_department__short_name__iexact=dept_param) |
                    Q(home_department__code__iexact=dept_param) |
                    Q(section__batch__course__department__short_name__iexact=dept_param) |
                    Q(section__batch__course__department__name__iexact=dept_param) |
                    Q(section__batch__department__short_name__iexact=dept_param)
                )
            if batch_param and batch_param != 'ALL':
                qs = qs.filter(section__batch__name__iexact=batch_param)
            if sec_param and sec_param != 'ALL':
                qs = qs.filter(section__name__iexact=sec_param)
            if search:
                clean_digits = ''.join([c for c in search if c.isdigit()])
                search_filters = (
                    Q(reg_no__icontains=search) |
                    Q(user__username__icontains=search) |
                    Q(user__first_name__icontains=search) |
                    Q(user__last_name__icontains=search) |
                    Q(user__email__icontains=search)
                )
                if clean_digits:
                    search_filters |= Q(reg_no__icontains=clean_digits)
                qs = qs.filter(search_filters)

        # Create in-memory zip
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            added_filenames = set()

            for st in qs:
                reg_no = str(st.reg_no).strip()
                # Use clean digits if digits exist, else exact string (identical to frontend extractBarcodeValue)
                clean_digits = ''.join([c for c in reg_no if c.isdigit()])
                barcode_str = clean_digits if (clean_digits and len(clean_digits) > 0) else reg_no

                # File name: <register_number>.png
                safe_reg_no = re.sub(r'[^a-zA-Z0-9_-]', '_', reg_no) or f"student_{st.id}"
                filename = f"{safe_reg_no}.png"

                # Avoid duplicate names in zip
                counter = 1
                while filename in added_filenames:
                    filename = f"{safe_reg_no}_{counter}.png"
                    counter += 1
                added_filenames.add(filename)

                out_img_bytes = None

                # Method 1: ReportLab Code128 barcode generator
                try:
                    barcode_drawing = code128.Code128(
                        barcode_str,
                        barHeight=25 * mm,
                        barWidth=0.5 * mm,
                        humanReadable=False
                    )
                    raw_img_buf = io.BytesIO()
                    from reportlab.graphics import renderPM
                    renderPM.drawToFile(barcode_drawing, raw_img_buf, fmt="PNG", dpi=300)
                    raw_img_buf.seek(0)

                    # Crop tightly using Pillow
                    pil_img = PILImage.open(raw_img_buf).convert("RGBA")
                    bg = PILImage.new("RGBA", pil_img.size, (255, 255, 255, 255))
                    from PIL import ImageChops
                    diff = ImageChops.difference(pil_img, bg)
                    bbox = diff.getbbox()

                    if bbox:
                        pad = 6
                        left = max(0, bbox[0] - pad)
                        top = max(0, bbox[1] - pad)
                        right = min(pil_img.width, bbox[2] + pad)
                        bottom = min(pil_img.height, bbox[3] + pad)
                        cropped = pil_img.crop((left, top, right, bottom))
                    else:
                        cropped = pil_img

                    out_buf = io.BytesIO()
                    cropped.save(out_buf, format="PNG")
                    out_img_bytes = out_buf.getvalue()
                except Exception as ex1:
                    # Method 2: Standard Code128 encoding via code128 writer
                    try:
                        # Direct Reportlab Drawing wrapper
                        from reportlab.graphics.shapes import Drawing
                        drawing = Drawing(180, 50)
                        bc = code128.Code128(clean_digits, barHeight=35, barWidth=1.2, humanReadable=False)
                        drawing.add(bc)
                        raw_buf = io.BytesIO()
                        from reportlab.graphics import renderPM
                        renderPM.drawToFile(drawing, raw_buf, fmt="PNG", dpi=300)
                        out_img_bytes = raw_buf.getvalue()
                    except Exception as ex2:
                        out_img_bytes = None

                if out_img_bytes:
                    zip_file.writestr(filename, out_img_bytes)

        zip_buffer.seek(0)
        response = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
        response['Content-Disposition'] = 'attachment; filename="student_barcodes_images.zip"'
        return response

