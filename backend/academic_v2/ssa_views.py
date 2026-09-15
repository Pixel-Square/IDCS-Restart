import logging
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.http import FileResponse
from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from .models import AcV2ExamAssignment, AcV2Section, AcV2SSAAssignment, AcV2SSASubmission
from .views import (
    _has_admin_bypass_access,
    _get_student_profile_for_request,
    _student_can_access_ta_ids,
    _course_header_from_section,
    _get_active_students_for_teaching_assignment,
    _save_draft_marks_for_exam,
)
from .ssa_utils import create_student_submission_pdf

logger = logging.getLogger(__name__)


def _ssa_exam_info(request, exam_assignment):
    from django.utils import timezone
    section = exam_assignment.section
    header = _course_header_from_section(section)
    
    edit_window_active = False
    if exam_assignment.edit_window_until and exam_assignment.edit_window_until > timezone.now():
        edit_window_active = True
        
    return {
        'exam_id': str(exam_assignment.id),
        'exam_type': str(exam_assignment.exam),
        'course_code': header['course_code'],
        'course_name': header['course_name'],
        'faculty_name': header['faculty_name'],
        'section_name': str(section),
        'semester': getattr(getattr(section, 'semester', None), 'semester_number', None),
        'max_marks': exam_assignment.max_marks,
        'has_pending_edit_request': exam_assignment.has_pending_edit_request,
        'edit_window_active': edit_window_active,
    }


def _ssa_assignment_payload(request, ssa_assignment):
    return {
        'id': str(ssa_assignment.id),
        'assignment_type': ssa_assignment.assignment_type,
        'status': ssa_assignment.status,
        'rubric_file_url': request.build_absolute_uri(ssa_assignment.rubric_file.url) if ssa_assignment.rubric_file else None,
        'first_page_file_url': None,  # Auto-generated, no manual upload needed
        'student_topic_file_url': request.build_absolute_uri(ssa_assignment.student_topic_file.url) if ssa_assignment.student_topic_file else None,
        'finalized_at': ssa_assignment.finalized_at,
    }

# --- Faculty Views ---

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def ssa_assignment_detail(request, exam_id):
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    
    # Auth
    if exam_assignment.section.faculty_user != request.user and not _has_admin_bypass_access(request.user):
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    if request.method == 'GET':
        try:
            ssa_assignment = exam_assignment.ssa_assignment
        except AcV2SSAAssignment.DoesNotExist:
            return Response({'exam_info': _ssa_exam_info(request, exam_assignment), 'assignment': None}, status=status.HTTP_200_OK)

        return Response({
            'exam_info': _ssa_exam_info(request, exam_assignment),
            'assignment': _ssa_assignment_payload(request, ssa_assignment),
        }, status=status.HTTP_200_OK)
        
    elif request.method == 'POST':
        raw_exam = str(exam_assignment.exam)
        assignment_type = raw_exam.replace(' ', '').replace('_', '').replace('-', '').upper()
        if assignment_type not in ['SSA1', 'SSA2']:
            return Response({'detail': f'Invalid assignment type {raw_exam}. Must be SSA1 or SSA2.'}, status=status.HTTP_400_BAD_REQUEST)
            
        try:
            ssa_assignment = exam_assignment.ssa_assignment
        except AcV2SSAAssignment.DoesNotExist:
            ssa_assignment = AcV2SSAAssignment(exam_assignment=exam_assignment, assignment_type=assignment_type)
            
        rubric_file = request.FILES.get('rubric_file') or request.FILES.get('rubrics')
        first_page_file = request.FILES.get('first_page_file') or request.FILES.get('first_page')
        student_topic_file = request.FILES.get('student_topic_file') or request.FILES.get('topic')
        if rubric_file:
            ssa_assignment.rubric_file = rubric_file
        if first_page_file:
            ssa_assignment.first_page_file = first_page_file
        if student_topic_file:
            if not student_topic_file.name.lower().endswith('.xlsx'):
                return Response({'detail': 'Student topic file must be an Excel template (.xlsx)'}, status=status.HTTP_400_BAD_REQUEST)
            ssa_assignment.student_topic_file = student_topic_file
            
        ssa_assignment.save()
        
        return Response({'detail': 'Saved successfully', 'id': ssa_assignment.id}, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def ssa_assignment_finalize(request, exam_id):
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    
    if exam_assignment.section.faculty_user != request.user and not _has_admin_bypass_access(request.user):
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    try:
        ssa_assignment = exam_assignment.ssa_assignment
    except AcV2SSAAssignment.DoesNotExist:
        return Response({'detail': 'SSA assignment not found'}, status=status.HTTP_404_NOT_FOUND)
        
    if not (ssa_assignment.rubric_file and ssa_assignment.student_topic_file):
        return Response({'detail': 'Rubric and Student Topic files are required to finalize.'}, status=status.HTTP_400_BAD_REQUEST)
        
    ssa_assignment.status = 'FINALIZED'
    ssa_assignment.finalized_at = timezone.now()
    ssa_assignment.finalized_by = request.user
    ssa_assignment.save(update_fields=['status', 'finalized_at', 'finalized_by'])
    
    return Response({'detail': 'Finalized successfully'}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_assignment_submissions(request, exam_id):
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    
    if exam_assignment.section.faculty_user != request.user and not _has_admin_bypass_access(request.user):
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    try:
        ssa_assignment = exam_assignment.ssa_assignment
    except AcV2SSAAssignment.DoesNotExist:
        return Response({'detail': 'SSA assignment not found'}, status=status.HTTP_404_NOT_FOUND)
        
    ta = exam_assignment.section.teaching_assignment
    students = _get_active_students_for_teaching_assignment(ta)
    
    submissions = {sub.student_id: sub for sub in AcV2SSASubmission.objects.filter(ssa_assignment=ssa_assignment)}
    
    result = []
    for sp in students:
        sub = submissions.get(sp.id)
        if sub:
            result.append({
                'student_id': sp.id,
                'student_name': str(sp.user) if sp.user else sp.reg_no,
                'reg_no': sp.reg_no,
                'submission_status': sub.submission_status,
                'submitted_at': sub.submitted_at,
                'marks': sub.marks,
                'evaluated_at': sub.evaluated_at,
                'submission_id': sub.id
            })
        else:
            result.append({
                'student_id': sp.id,
                'student_name': str(sp.user) if sp.user else sp.reg_no,
                'reg_no': sp.reg_no,
                'submission_status': 'NOT_SUBMITTED',
                'submitted_at': None,
                'marks': None,
                'evaluated_at': None,
                'submission_id': None
            })
            
    return Response({
        'assignment_info': {
            **_ssa_exam_info(request, exam_assignment),
            'status': ssa_assignment.status,
        },
        'submissions': result,
    }, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_submission_detail(request, submission_id):
    submission = get_object_or_404(AcV2SSASubmission, id=submission_id)
    exam_assignment = submission.ssa_assignment.exam_assignment
    
    if exam_assignment.section.faculty_user != request.user and not _has_admin_bypass_access(request.user):
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    all_submissions = list(AcV2SSASubmission.objects.filter(
        ssa_assignment=submission.ssa_assignment,
    ).order_by('student__reg_no'))
    return Response({
        'submission': {
            'id': str(submission.id),
            'submission_id': str(submission.id),
            'student_id': submission.student_id,
            'student_name': submission.student_name,
            'reg_no': submission.reg_no,
            'submission_status': submission.submission_status,
            'submitted_at': submission.submitted_at,
            'marks': submission.marks,
            'feedback': submission.feedback,
            'evaluated_at': submission.evaluated_at,
            'submitted_file_url': request.build_absolute_uri(submission.submitted_file.url) if getattr(submission, 'submitted_file', None) and submission.submitted_file.name else None,
            'generated_file_url': request.build_absolute_uri(submission.generated_file.url) if getattr(submission, 'generated_file', None) and submission.generated_file.name else (request.build_absolute_uri(submission.submitted_file.url) if getattr(submission, 'submitted_file', None) and submission.submitted_file.name else None),
        },
        'assignment_info': _ssa_exam_info(request, exam_assignment),
        'all_submissions': [
            {'submission_id': str(item.id), 'student_name': item.student_name, 'reg_no': item.reg_no}
            for item in all_submissions
        ],
    }, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def ssa_submission_evaluate(request, submission_id):
    submission = get_object_or_404(AcV2SSASubmission, id=submission_id)
    exam_assignment = submission.ssa_assignment.exam_assignment
    
    if exam_assignment.section.faculty_user != request.user and not _has_admin_bypass_access(request.user):
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    from django.utils import timezone
    edit_window_active = bool(exam_assignment.edit_window_until and exam_assignment.edit_window_until > timezone.now())
    if submission.submission_status == 'EVALUATED' and not edit_window_active:
        return Response({'detail': 'Evaluation locked. Please request edit access.'}, status=status.HTTP_403_FORBIDDEN)
        
    marks_value = request.data.get('marks')
    feedback = request.data.get('feedback', '')
    
    if marks_value is None:
        return Response({'detail': 'Marks required'}, status=status.HTTP_400_BAD_REQUEST)
        
    try:
        marks_float = float(marks_value)
    except ValueError:
        return Response({'detail': 'Invalid marks'}, status=status.HTTP_400_BAD_REQUEST)
        
    if marks_float < 0 or marks_float > float(exam_assignment.max_marks):
        return Response({'detail': f'Marks must be between 0 and {exam_assignment.max_marks}'}, status=status.HTTP_400_BAD_REQUEST)
        
    submission.marks = marks_float
    submission.feedback = feedback
    submission.evaluated_by = request.user
    submission.evaluated_at = timezone.now()
    submission.submission_status = 'EVALUATED'
    submission.save()
    
    marks_payload = [{
        'student_id': str(submission.student_id),
        'mark': marks_float,
        'co_marks': {},
        'is_absent': False,
    }]
    _save_draft_marks_for_exam(exam_assignment, marks_payload)
    
    draft = exam_assignment.draft_data if isinstance(exam_assignment.draft_data, dict) else {}
    marks_map = draft.get('marks', {}) if isinstance(draft.get('marks', {}), dict) else {}
    marks_map[str(submission.student_id)] = {'mark': marks_float, 'co_marks': {}, 'is_absent': False}
    draft['marks'] = marks_map
    exam_assignment.draft_data = draft
    exam_assignment.last_saved_at = timezone.now()
    exam_assignment.last_saved_by = request.user
    exam_assignment.save(update_fields=['draft_data', 'last_saved_at', 'last_saved_by'])
    
    return Response({'detail': 'Evaluation saved'}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_assignment_topic_template(request, exam_id):
    from academic_v2.views import _get_active_students_for_teaching_assignment
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.http import HttpResponse
    
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    if exam_assignment.section.faculty_user != request.user and not _has_admin_bypass_access(request.user):
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    try:
        num_topics = int(request.GET.get('num_topics', 1))
        if num_topics < 1:
            num_topics = 1
    except ValueError:
        num_topics = 1
        
    ta = exam_assignment.section.teaching_assignment
    active_student_profiles = _get_active_students_for_teaching_assignment(ta)
    # Sort students by registration number
    active_student_profiles = sorted(active_student_profiles, key=lambda sp: sp.reg_no or '')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Topics"
    
    # Headers
    headers = ["Reg No", "Student Name"]
    for i in range(1, num_topics + 1):
        headers.append(f"Topic {i}")
        
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Adjust column widths
        if col_idx == 1:
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
        elif col_idx == 2:
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 35
        else:
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 40
            
    # Rows
    for row_idx, sp in enumerate(active_student_profiles, 2):
        ws.cell(row=row_idx, column=1, value=sp.reg_no or '')
        ws.cell(row=row_idx, column=2, value=str(sp.user) if sp.user else '')
        
    from io import BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    response = HttpResponse(buffer.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="student_topics_template.xlsx"'
    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_assignment_file(request, exam_id, file_type):
    """Serve SSA assignment files (rubric, first_page, topic) to authorized faculty or students."""
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    try:
        ssa_assignment = exam_assignment.ssa_assignment
    except AcV2SSAAssignment.DoesNotExist:
        return Response({'detail': 'Not found'}, status=status.HTTP_404_NOT_FOUND)

    is_faculty = (exam_assignment.section.faculty_user == request.user or _has_admin_bypass_access(request.user))
    sp = _get_student_profile_for_request(request)
    is_student = False
    if sp:
        allowed_tas = _student_can_access_ta_ids(sp)
        is_student = exam_assignment.section.teaching_assignment_id in allowed_tas

    if not is_faculty and not is_student:
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)

    file_map = {
        'rubric': ssa_assignment.rubric_file,
        'first_page': ssa_assignment.first_page_file,
        'topic': ssa_assignment.student_topic_file,
    }
    file_field = file_map.get(file_type)
    if not file_field:
        return Response({'detail': 'File not found'}, status=status.HTTP_404_NOT_FOUND)

    try:
        return FileResponse(file_field.open('rb'), as_attachment=True, filename=file_field.name.split('/')[-1])
    except IOError:
        return Response({'detail': 'File unavailable'}, status=status.HTTP_404_NOT_FOUND)


# --- Student Views ---

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_student_assignments(request):
    sp = _get_student_profile_for_request(request)
    if not sp:
        return Response({'detail': 'Student profile not found'}, status=status.HTTP_403_FORBIDDEN)
        
    allowed_tas = _student_can_access_ta_ids(sp)
    
    assignments = AcV2SSAAssignment.objects.filter(
        status='FINALIZED',
        exam_assignment__section__teaching_assignment_id__in=allowed_tas
    ).select_related('exam_assignment__section__course', 'exam_assignment__section__faculty_user')
    
    submissions = {sub.ssa_assignment_id: sub for sub in AcV2SSASubmission.objects.filter(student=sp)}
    
    result = []
    for ssa in assignments:
        sec = ssa.exam_assignment.section
        header = _course_header_from_section(sec)
        sub = submissions.get(ssa.id)
        
        result.append({
            'exam_id': ssa.exam_assignment.id,
            'assignment_type': ssa.assignment_type,
            'course_code': header['course_code'],
            'course_name': header['course_name'],
            'faculty_name': header['faculty_name'],
            'submission_status': sub.submission_status if sub else 'NOT_SUBMITTED'
        })
    return Response({'assignments': result}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_student_assignment_detail(request, exam_id):
    sp = _get_student_profile_for_request(request)
    if not sp:
        return Response({'detail': 'Student profile not found'}, status=status.HTTP_403_FORBIDDEN)
        
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    allowed_tas = _student_can_access_ta_ids(sp)
    if exam_assignment.section.teaching_assignment_id not in allowed_tas:
        return Response({'detail': 'Not enrolled'}, status=status.HTTP_403_FORBIDDEN)
        
    try:
        ssa_assignment = exam_assignment.ssa_assignment
    except AcV2SSAAssignment.DoesNotExist:
        return Response({'detail': 'Not found'}, status=status.HTTP_404_NOT_FOUND)
        
    sub = AcV2SSASubmission.objects.filter(ssa_assignment=ssa_assignment, student=sp).first()
    header = _course_header_from_section(exam_assignment.section)
    
    return Response({
        'exam_id': exam_assignment.id,
        'assignment_type': ssa_assignment.assignment_type,
        'course_code': header['course_code'],
        'course_name': header['course_name'],
        'faculty_name': header['faculty_name'],
        'rubric_file_url': request.build_absolute_uri(ssa_assignment.rubric_file.url) if ssa_assignment.rubric_file and ssa_assignment.rubric_file.name else None,
        'first_page_file_url': request.build_absolute_uri(ssa_assignment.first_page_file.url) if ssa_assignment.first_page_file and ssa_assignment.first_page_file.name else None,
        'student_topic_file_url': request.build_absolute_uri(ssa_assignment.student_topic_file.url) if ssa_assignment.student_topic_file and ssa_assignment.student_topic_file.name else None,
        'submission_status': sub.submission_status if sub else 'NOT_SUBMITTED',
        'marks': sub.marks if sub else None,
        'feedback': sub.feedback if sub else '',
        'submission_id': sub.id if sub else None,
        'submitted_at': sub.submitted_at if sub else None,
        'generated_file_url': request.build_absolute_uri(sub.generated_file.url) if sub and getattr(sub, 'generated_file', None) and sub.generated_file.name else (request.build_absolute_uri(sub.submitted_file.url) if sub and getattr(sub, 'submitted_file', None) and sub.submitted_file.name else None),
    }, status=status.HTTP_200_OK)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def ssa_student_submit(request, exam_id):
    sp = _get_student_profile_for_request(request)
    if not sp:
        return Response({'detail': 'Student profile not found'}, status=status.HTTP_403_FORBIDDEN)
        
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    allowed_tas = _student_can_access_ta_ids(sp)
    if exam_assignment.section.teaching_assignment_id not in allowed_tas:
        return Response({'detail': 'Not enrolled'}, status=status.HTTP_403_FORBIDDEN)
        
    try:
        ssa_assignment = exam_assignment.ssa_assignment
    except AcV2SSAAssignment.DoesNotExist:
        return Response({'detail': 'Not found'}, status=status.HTTP_404_NOT_FOUND)
        
    if 'file' not in request.FILES:
        return Response({'detail': 'File required'}, status=status.HTTP_400_BAD_REQUEST)
        
    upload_file = request.FILES['file']
    if not upload_file.name.lower().endswith('.pdf') or upload_file.content_type != 'application/pdf':
        return Response({'detail': 'Only PDF files are allowed'}, status=status.HTTP_400_BAD_REQUEST)
        
    sub, created = AcV2SSASubmission.objects.get_or_create(
        ssa_assignment=ssa_assignment,
        student=sp,
        defaults={
            'reg_no': sp.reg_no or '',
            'student_name': str(sp.user) if sp.user else (sp.reg_no or '')
        }
    )
    
    if not created and sub.submission_status in ['UNDER_EVALUATION', 'EVALUATED']:
        return Response({'detail': 'Cannot resubmit after evaluation has started'}, status=status.HTTP_400_BAD_REQUEST)
        
    sub.submitted_file = upload_file
    sub.save()
    
    try:
        from .ssa_utils import create_student_submission_pdf, get_student_topics_from_excel
        topics = get_student_topics_from_excel(ssa_assignment, sp.reg_no or '')
        generated_file = create_student_submission_pdf(sp, ssa_assignment, upload_file, topics=topics)
        if generated_file:
            sub.generated_file.save(generated_file.name, generated_file, save=False)
    except Exception as e:
        logger.error(f"Error generating PDF for submission {sub.id}: {e}")
        
    sub.submission_status = 'SUBMITTED'
    sub.submitted_at = timezone.now()
    sub.save()
    
    return Response({'detail': 'Submitted successfully'}, status=status.HTTP_200_OK)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_student_submission_file(request, submission_id):
    sp = _get_student_profile_for_request(request)
    if not sp:
        return Response({'detail': 'Student profile not found'}, status=status.HTTP_403_FORBIDDEN)
        
    submission = get_object_or_404(AcV2SSASubmission, id=submission_id)
    if submission.student_id != sp.id:
        return Response({'detail': 'Not authorized'}, status=status.HTTP_403_FORBIDDEN)
        
    if not submission.generated_file:
        if not submission.submitted_file:
            return Response({'detail': 'File not found'}, status=status.HTTP_404_NOT_FOUND)
        file_field = submission.submitted_file
    else:
        file_field = submission.generated_file
        
    try:
        return FileResponse(file_field.open('rb'), as_attachment=True, filename=file_field.name.split('/')[-1])
    except IOError:
        return Response({'detail': 'File unavailable'}, status=status.HTTP_404_NOT_FOUND)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def ssa_student_submission_status(request, exam_id):
    sp = _get_student_profile_for_request(request)
    if not sp:
        return Response({'detail': 'Student profile not found'}, status=status.HTTP_403_FORBIDDEN)
        
    exam_assignment = get_object_or_404(AcV2ExamAssignment, id=exam_id)
    allowed_tas = _student_can_access_ta_ids(sp)
    if exam_assignment.section.teaching_assignment_id not in allowed_tas:
        return Response({'detail': 'Not enrolled'}, status=status.HTTP_403_FORBIDDEN)
        
    try:
        ssa_assignment = exam_assignment.ssa_assignment
    except AcV2SSAAssignment.DoesNotExist:
        return Response({'detail': 'Not found'}, status=status.HTTP_404_NOT_FOUND)
        
    sub = AcV2SSASubmission.objects.filter(ssa_assignment=ssa_assignment, student=sp).first()
    if not sub:
        return Response({'status': 'NOT_SUBMITTED'}, status=status.HTTP_200_OK)
        
    return Response({
        'submission_status': sub.submission_status,
        'marks': sub.marks,
        'feedback': sub.feedback,
        'submitted_at': sub.submitted_at,
        'evaluated_at': sub.evaluated_at
    }, status=status.HTTP_200_OK)
